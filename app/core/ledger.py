"""报销台账导出：CSV / XLSX / 可打印 HTML。"""

import csv
import io
import logging
from typing import Any, Dict, List

logger = logging.getLogger(__name__)

COLUMNS = [
    ("序号", None),
    ("发票代码", ("通用字段", "发票代码")),
    ("发票号码", ("通用字段", "发票号码")),
    ("开票日期", ("通用字段", "开票日期")),
    ("购买方", ("购买方信息", "名称")),
    ("销售方", ("销售方信息", "名称")),
    ("商品名称", ("商品信息", "商品名称")),
    ("金额", ("通用字段", "金额")),
    ("税额", ("通用字段", "税额")),
    ("价税合计", ("通用字段", "价税合计")),
]


def _pick(fields: Dict[str, Any], path) -> str:
    if path is None:
        return ""
    cat, key = path
    value = (fields.get(cat) or {}).get(key, "")
    return "" if value in (None, "无") else str(value)


def build_rows(records: List[Dict[str, Any]]) -> List[List[str]]:
    rows = []
    for i, rec in enumerate(records, 1):
        fields = rec.get("fields", {})
        rows.append([str(i)] + [_pick(fields, path) for _, path in COLUMNS[1:]])
    return rows


def to_csv(records: List[Dict[str, Any]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c[0] for c in COLUMNS])
    writer.writerows(build_rows(records))
    return "\ufeff" + buf.getvalue()  # 带 BOM，Excel 中文不乱码


def to_xlsx(records: List[Dict[str, Any]], title: str = "报销台账") -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("未安装 openpyxl，无法导出 XLSX。请 `pip install openpyxl`。") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "报销台账"

    header = [c[0] for c in COLUMNS]
    ws.append(header)
    fill = PatternFill("solid", fgColor="2563EB")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for row in build_rows(records):
        ws.append(row)

    widths = [6, 16, 16, 14, 24, 24, 20, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_html(records: List[Dict[str, Any]], title: str = "报销单") -> str:
    total = 0.0
    for rec in records:
        raw = (rec.get("fields", {}).get("通用字段") or {}).get("价税合计", "")
        try:
            total += float(str(raw).replace(",", ""))
        except (ValueError, TypeError):
            pass

    head = "".join(f"<th>{c[0]}</th>" for c in COLUMNS)
    body = "".join(
        "<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>"
        for row in build_rows(records)
    )
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="UTF-8"><title>{title}</title>
<style>
  body {{ font-family: "Microsoft YaHei", sans-serif; margin: 32px; color: #1f2937; }}
  h1 {{ font-size: 20px; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
  th, td {{ border: 1px solid #cbd5e1; padding: 6px 8px; text-align: left; }}
  th {{ background: #f1f5f9; }}
  tfoot td {{ font-weight: bold; }}
  @media print {{ body {{ margin: 0; }} }}
</style></head><body>
<h1>{title}</h1>
<p>发票张数：{len(records)}　合计金额：{total:.2f} 元</p>
<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody>
<tfoot><tr><td colspan="{len(COLUMNS)-1}" style="text-align:right">合计</td><td>{total:.2f}</td></tr></tfoot>
</table>
<p style="margin-top:24px;font-size:12px;color:#64748b">由「小荷包报销助手」生成</p>
</body></html>"""
